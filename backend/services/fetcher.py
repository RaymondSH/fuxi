"""内容抓取：把各种来源统一抓成「标题 + 正文文本」。

支持的来源：
- url    网页 / 博客（httpx）；微信公众号走 Jina Reader
- pdf    PDF 文件（pypdf）
- docx   Word 文件（python-docx）
- xlsx   Excel 文件（openpyxl，转成 Markdown 表格）
- image  图片（交给 GLM 视觉模型描述）

每个抓取函数返回 FetchResult，后续 ingest_worker 拿 content 去做摘要/要点/实体。
"""
from __future__ import annotations

import base64
import io
import re
from dataclasses import dataclass

import httpx

from config import settings
from services import llm

# 图片扩展名 → media_type，GLM 视觉接口需要
_IMAGE_MEDIA = {
    "png": "image/png",
    "jpg": "image/jpeg",
    "jpeg": "image/jpeg",
    "gif": "image/gif",
    "webp": "image/webp",
}

_HEADERS = {
    "User-Agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
    "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0 Safari/537.36"
}


@dataclass
class FetchResult:
    title: str
    content: str  # 纯文本正文
    source_type: str  # 'url' | 'pdf' | 'docx' | 'xlsx' | 'image'
    raw_bytes: bytes | None = None  # 原始文件字节，供上传对象存储


# ---------- URL ----------

def fetch_url(url: str) -> FetchResult:
    """抓取网页。微信公众号反爬严重，改用 Jina Reader 拿干净 Markdown。"""
    if "mp.weixin.qq.com" in url:
        return _fetch_via_jina(url)

    with httpx.Client(headers=_HEADERS, follow_redirects=True, timeout=30) as c:
        resp = c.get(url)
        resp.raise_for_status()
        html = resp.text

    title = _extract_title(html)
    content = _html_to_text(html)
    return FetchResult(title=title, content=content, source_type="url")


def _fetch_via_jina(url: str) -> FetchResult:
    """Jina Reader：在原链接前拼前缀，返回提取好的 Markdown。"""
    with httpx.Client(headers=_HEADERS, follow_redirects=True, timeout=60) as c:
        resp = c.get(settings.jina_reader_prefix + url)
        resp.raise_for_status()
        md = resp.text
    # Jina 输出头部带 "Title: xxx"，提取出来当标题
    title = "未命名"
    m = re.search(r"^Title:\s*(.+)$", md, re.MULTILINE)
    if m:
        title = m.group(1).strip()
    return FetchResult(title=title, content=md, source_type="url")


_TITLE_RE = re.compile(r"<title[^>]*>(.*?)</title>", re.IGNORECASE | re.DOTALL)
_SCRIPT_STYLE_RE = re.compile(r"<(script|style)[^>]*>.*?</\1>", re.IGNORECASE | re.DOTALL)
_TAG_RE = re.compile(r"<[^>]+>")
_WS_RE = re.compile(r"\n\s*\n\s*\n+")


def _extract_title(html: str) -> str:
    m = _TITLE_RE.search(html)
    return m.group(1).strip() if m else "未命名"


def _html_to_text(html: str) -> str:
    """轻量正文提取：去脚本/样式/标签。够 MVP 用；要更干净可换 readability/trafilatura。"""
    html = _SCRIPT_STYLE_RE.sub(" ", html)
    html = re.sub(r"<br\s*/?>", "\n", html, flags=re.IGNORECASE)
    html = re.sub(r"</(p|div|h[1-6]|li)>", "\n", html, flags=re.IGNORECASE)
    text = _TAG_RE.sub("", html)
    text = _unescape(text)
    return _WS_RE.sub("\n\n", text).strip()


def _unescape(text: str) -> str:
    import html as _h

    return _h.unescape(text)


# ---------- PDF ----------

def fetch_pdf(data: bytes, *, filename: str = "") -> FetchResult:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = [page.extract_text() or "" for page in reader.pages]
    content = "\n\n".join(pages).strip()
    title = (reader.metadata.title if reader.metadata else None) or _stem(filename) or "未命名 PDF"
    return FetchResult(title=title, content=content, source_type="pdf", raw_bytes=data)


# ---------- Word ----------

def fetch_docx(data: bytes, *, filename: str = "") -> FetchResult:
    from docx import Document

    doc = Document(io.BytesIO(data))
    paras = [p.text for p in doc.paragraphs if p.text.strip()]
    content = "\n\n".join(paras).strip()
    # Word 第一段常是标题
    title = (doc.core_properties.title or (paras[0] if paras else "") or _stem(filename) or "未命名文档")
    return FetchResult(title=title, content=content, source_type="docx", raw_bytes=data)


# ---------- Excel ----------

def fetch_xlsx(data: bytes, *, filename: str = "") -> FetchResult:
    """每个 sheet 转成 Markdown 表格，方便 LLM 理解结构。"""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    blocks: list[str] = []
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        blocks.append(f"## {ws.title}\n\n" + _rows_to_markdown(rows))
    content = "\n\n".join(blocks).strip()
    title = _stem(filename) or "未命名表格"
    return FetchResult(title=title, content=content, source_type="xlsx", raw_bytes=data)


def _rows_to_markdown(rows: list[tuple]) -> str:
    def fmt(row: tuple) -> str:
        return "| " + " | ".join("" if v is None else str(v) for v in row) + " |"

    header, *body = rows
    sep = "| " + " | ".join("---" for _ in header) + " |"
    lines = [fmt(header), sep, *(fmt(r) for r in body)]
    return "\n".join(lines)


# ---------- 图片 ----------

def fetch_image(data: bytes, *, filename: str = "") -> FetchResult:
    """图片没有「文本」，交给 GLM 视觉模型转写成可入库的文字。"""
    ext = _ext(filename)
    media_type = _IMAGE_MEDIA.get(ext, "image/png")
    b64 = base64.standard_b64encode(data).decode("utf-8")
    content = llm.describe_image(b64, media_type)
    title = _stem(filename) or "未命名图片"
    return FetchResult(title=title, content=content, source_type="image", raw_bytes=data)


# ---------- 统一入口 ----------

def fetch(*, url: str | None = None, data: bytes | None = None, filename: str = "") -> FetchResult:
    """按输入类型分派。给 url 走网页；给文件字节按扩展名分派。"""
    if url:
        return fetch_url(url)
    if data is None:
        raise ValueError("必须提供 url 或文件 data")

    ext = _ext(filename)
    if ext == "pdf":
        return fetch_pdf(data, filename=filename)
    if ext == "docx":
        return fetch_docx(data, filename=filename)
    if ext in ("xlsx", "xls"):
        return fetch_xlsx(data, filename=filename)
    if ext in _IMAGE_MEDIA:
        return fetch_image(data, filename=filename)
    raise ValueError(f"不支持的文件类型：{filename!r}")


# ---------- 小工具 ----------

def _ext(filename: str) -> str:
    return filename.rsplit(".", 1)[-1].lower() if "." in filename else ""


def _stem(filename: str) -> str:
    base = filename.rsplit("/", 1)[-1]
    return base.rsplit(".", 1)[0] if "." in base else base
